import openpyxl

def getRowCount(file, sheetName):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetName]
    return (sheet.max_row)

def getColumnCount(file,Sheetname):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[Sheetname]
    return (sheet.max_column)

def readData(file,sheetName,rowNum,columnno):
    workbook = openpyxl.load_workbook(file)
    sheet =  workbook[sheetName]
    return sheet.cell(row=rowNum,column=columnno).value

def writeData(file,sheetname,rownum,columnno,data):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetname]
    sheet.cell(row=rownum,column=columnno).value = data
    workbook.save(file)
